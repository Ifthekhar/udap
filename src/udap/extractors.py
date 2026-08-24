"""Input extraction boundary for DOCX and PDF sources.

The first implementation keeps third-party imports optional so the core rule
engine can run before local dependencies are installed.
"""

from __future__ import annotations

from pathlib import Path
from statistics import median
from typing import Any

from .models import DocumentElement, DocumentModel, ElementType, SourceLocation
from .pdf_inspection import PdfInspectionDependencyError, inspect_pdf


class UnsupportedDocumentError(ValueError):
    pass


class MissingExtractorDependencyError(RuntimeError):
    pass


def load_document(path: str | Path) -> DocumentModel:
    source = Path(path)
    suffix = source.suffix.lower()

    if suffix == ".docx":
        return load_docx(source)
    if suffix == ".pdf":
        return load_pdf(source)

    raise UnsupportedDocumentError(f"Unsupported document format: {suffix or 'unknown'}")


def load_docx(path: Path) -> DocumentModel:
    try:
        import docx  # type: ignore[import-not-found]
    except ImportError as exc:
        raise MissingExtractorDependencyError(
            "DOCX extraction requires python-docx. Install project dependencies first."
        ) from exc

    document = docx.Document(path)
    elements: list[DocumentElement] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.lower().startswith("heading"):
            level = _parse_heading_level(style_name)
            elements.append(
                DocumentElement(
                    type=ElementType.HEADING,
                    text=text,
                    heading_level=level,
                    source=SourceLocation(element_id=str(index), description=style_name),
                )
            )
        else:
            elements.append(
                DocumentElement(
                    type=ElementType.PARAGRAPH,
                    text=text,
                    source=SourceLocation(element_id=str(index), description=style_name),
                )
            )

    core = document.core_properties
    return DocumentModel(
        original_filename=path.name,
        source_format="docx",
        title=core.title or None,
        language=None,
        elements=elements,
        metadata={"author": core.author, "subject": core.subject},
    )


def load_pdf(path: Path) -> DocumentModel:
    try:
        import pymupdf as fitz  # type: ignore[import-not-found]
    except ImportError as exc:
        try:
            import fitz  # type: ignore[import-not-found,no-redef]
        except ImportError:
            raise MissingExtractorDependencyError(
                "PDF extraction requires PyMuPDF. Install project dependencies first."
            ) from exc

    try:
        inspection = inspect_pdf(path)
    except PdfInspectionDependencyError as exc:
        raise MissingExtractorDependencyError(str(exc)) from exc

    pdf = fitz.open(path)
    elements: list[DocumentElement] = []

    for page_index, page in enumerate(pdf, start=1):
        text_blocks = _extract_page_text_blocks(page)
        body_font_size = _median_font_size(text_blocks)
        if _looks_multi_column(text_blocks):
            inspection.pages_with_multiple_columns.append(page_index)

        for block_index, block in enumerate(text_blocks, start=1):
            content = block["text"]
            if not content:
                continue

            element_type = ElementType.HEADING if _looks_like_heading(block, body_font_size) else ElementType.PARAGRAPH
            heading_level = _infer_heading_level(block, body_font_size) if element_type == ElementType.HEADING else None
            if element_type == ElementType.HEADING:
                inspection.heading_candidate_count += 1

            if _looks_like_table_text(content):
                element_type = ElementType.TABLE
                inspection.table_candidate_count += 1

            elements.append(
                DocumentElement(
                    type=element_type,
                    text=content,
                    source=SourceLocation(
                        page_number=page_index,
                        element_id=f"{page_index}:{block_index}",
                        bbox=block["bbox"],
                    ),
                    confidence=_block_confidence(block, text_blocks),
                    heading_level=heading_level,
                    table_headers=_infer_table_headers(content)
                    if element_type == ElementType.TABLE
                    else [],
                    metadata={
                        "font_size": block["font_size"],
                        "line_count": block["line_count"],
                        "origin": "pdf_text_block",
                    },
                )
            )

        image_rects = _extract_image_rects(page)
        for image_index, image in enumerate(page.get_images(full=True), start=1):
            bbox = image_rects[image_index - 1] if image_index <= len(image_rects) else None
            elements.append(
                DocumentElement(
                    type=ElementType.IMAGE,
                    source=SourceLocation(
                        page_number=page_index,
                        element_id=f"{page_index}:image:{image_index}",
                        description="PDF image object",
                        bbox=bbox,
                    ),
                    confidence=0.65,
                    metadata={
                        "xref": image[0] if image else None,
                        "nearby_text": _nearby_text(bbox, text_blocks) if bbox else "",
                    },
                )
            )

        for link_index, link in enumerate(page.get_links(), start=1):
            uri = link.get("uri")
            rect = _rect_to_bbox(link.get("from"))
            link_text = _text_overlapping_rect(rect, text_blocks) or uri or ""
            elements.append(
                DocumentElement(
                    type=ElementType.LINK,
                    text=link_text,
                    href=uri,
                    source=SourceLocation(
                        page_number=page_index,
                        element_id=f"{page_index}:link:{link_index}",
                        description="PDF link annotation",
                        bbox=rect,
                    ),
                    confidence=0.85 if link_text and uri else 0.45,
                    metadata={"kind": link.get("kind"), "nearby_text": _nearby_text(rect, text_blocks)},
                )
            )

    metadata = pdf.metadata or {}
    return DocumentModel(
        original_filename=path.name,
        source_format="pdf",
        title=inspection.title or metadata.get("title") or None,
        language=inspection.language,
        elements=elements,
        metadata=metadata,
        pdf=inspection,
    )


def _parse_heading_level(style_name: str) -> int | None:
    parts = style_name.split()
    if not parts:
        return None
    try:
        return int(parts[-1])
    except ValueError:
        return None


def _extract_block_text(block: object) -> str:
    if not isinstance(block, (list, tuple)) or len(block) < 5:
        return ""
    return " ".join(line.strip() for line in str(block[4]).splitlines() if line.strip())


def _extract_page_text_blocks(page: Any) -> list[dict[str, Any]]:
    raw = page.get_text("dict")
    blocks: list[dict[str, Any]] = []

    for block in raw.get("blocks", []):
        if block.get("type") != 0:
            continue

        lines = block.get("lines", [])
        for line in lines:
            line_text_parts: list[str] = []
            font_sizes: list[float] = []
            for span in line.get("spans", []):
                text = str(span.get("text", "")).strip()
                if not text:
                    continue
                line_text_parts.append(text)
                size = span.get("size")
                if isinstance(size, (int, float)):
                    font_sizes.append(float(size))

            line_text = " ".join(line_text_parts).strip()
            if not line_text:
                continue

            blocks.append(
                {
                    "bbox": _rect_to_bbox(line.get("bbox") or block.get("bbox")),
                    "text": line_text,
                    "font_size": max(font_sizes) if font_sizes else 0.0,
                    "line_count": 1,
                }
            )

    return sorted(blocks, key=lambda item: (item["bbox"][1], item["bbox"][0]))


def _median_font_size(blocks: list[dict[str, Any]]) -> float:
    sizes = [block["font_size"] for block in blocks if block["font_size"]]
    if not sizes:
        return 0.0
    return float(median(sizes))


def _looks_like_heading(block: dict[str, Any], body_font_size: float) -> bool:
    text = block["text"]
    if len(text) > 120 or block["line_count"] > 2:
        return False
    if body_font_size and block["font_size"] >= body_font_size + 2:
        return True
    return text.isupper() and len(text.split()) <= 10


def _infer_heading_level(block: dict[str, Any], body_font_size: float) -> int:
    if body_font_size and block["font_size"] >= body_font_size + 8:
        return 1
    if body_font_size and block["font_size"] >= body_font_size + 4:
        return 2
    return 3


def _looks_like_table_text(text: str) -> bool:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) < 2:
        return False
    delimiter_rows = sum(1 for line in lines if "\t" in line or "  " in line)
    return delimiter_rows >= 2


def _infer_table_headers(text: str) -> list[str]:
    first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")
    if not first_line:
        return []
    if "\t" in first_line:
        return [part.strip() for part in first_line.split("\t") if part.strip()]
    return [part.strip() for part in first_line.split("  ") if part.strip()]


def _block_confidence(block: dict[str, Any], blocks: list[dict[str, Any]]) -> float:
    if block["line_count"] == 1:
        return 0.85
    return 0.8


def _looks_multi_column(blocks: list[dict[str, Any]]) -> bool:
    if len(blocks) < 6:
        return False

    left_edges = [block["bbox"][0] for block in blocks if block["bbox"]]
    if len(left_edges) < 6:
        return False

    clusters: list[list[float]] = []
    for edge in sorted(left_edges):
        if not clusters or abs(edge - median(clusters[-1])) > 48:
            clusters.append([edge])
        else:
            clusters[-1].append(edge)

    substantial_clusters = [cluster for cluster in clusters if len(cluster) >= 2]
    if len(substantial_clusters) < 2:
        return False

    return median(substantial_clusters[-1]) - median(substantial_clusters[0]) > 120


def _extract_image_rects(page: Any) -> list[tuple[float, float, float, float]]:
    rects: list[tuple[float, float, float, float]] = []
    for image in page.get_images(full=True):
        xref = image[0] if image else None
        if xref is None:
            continue
        try:
            for rect in page.get_image_rects(xref):
                bbox = _rect_to_bbox(rect)
                if bbox:
                    rects.append(bbox)
        except RuntimeError:
            continue
    return rects


def _text_overlapping_rect(
    rect: tuple[float, float, float, float] | None,
    blocks: list[dict[str, Any]],
) -> str:
    if rect is None:
        return ""
    matches = [block["text"] for block in blocks if _rects_overlap(rect, block["bbox"])]
    return " ".join(matches).strip()


def _nearby_text(
    rect: tuple[float, float, float, float] | None,
    blocks: list[dict[str, Any]],
    limit: int = 160,
) -> str:
    if rect is None:
        return ""
    nearby = sorted(
        blocks,
        key=lambda block: _rect_distance(rect, block["bbox"]),
    )
    return " ".join(block["text"] for block in nearby[:2])[:limit].strip()


def _rect_to_bbox(rect: object) -> tuple[float, float, float, float] | None:
    if rect is None:
        return None
    if isinstance(rect, (list, tuple)) and len(rect) >= 4:
        return (float(rect[0]), float(rect[1]), float(rect[2]), float(rect[3]))
    if all(hasattr(rect, attr) for attr in ("x0", "y0", "x1", "y1")):
        return (float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1))
    return None


def _rects_overlap(
    first: tuple[float, float, float, float] | None,
    second: tuple[float, float, float, float] | None,
) -> bool:
    if first is None or second is None:
        return False
    return not (
        first[2] < second[0]
        or second[2] < first[0]
        or first[3] < second[1]
        or second[3] < first[1]
    )


def _rect_distance(
    first: tuple[float, float, float, float] | None,
    second: tuple[float, float, float, float] | None,
) -> float:
    if first is None or second is None:
        return float("inf")
    first_center = ((first[0] + first[2]) / 2, (first[1] + first[3]) / 2)
    second_center = ((second[0] + second[2]) / 2, (second[1] + second[3]) / 2)
    return abs(first_center[0] - second_center[0]) + abs(first_center[1] - second_center[1])
